"""
Script convert file Excel (Bảng tính calo thực phẩm.xlsx) thành JSON
Chạy: python tool/excel_to_json.py
Output: assets/database/food_database_from_excel.json
"""

import openpyxl
import json
import re
import unicodedata

def create_key(name: str) -> str:
    """Tạo key từ tên thực phẩm tiếng Việt"""
    # Bỏ dấu tiếng Việt
    text = name.lower().strip()
    
    # Mapping Vietnamese characters
    vietnamese_map = {
        'à': 'a', 'á': 'a', 'ả': 'a', 'ã': 'a', 'ạ': 'a',
        'ă': 'a', 'ắ': 'a', 'ằ': 'a', 'ẳ': 'a', 'ẵ': 'a', 'ặ': 'a',
        'â': 'a', 'ấ': 'a', 'ầ': 'a', 'ẩ': 'a', 'ẫ': 'a', 'ậ': 'a',
        'đ': 'd',
        'è': 'e', 'é': 'e', 'ẻ': 'e', 'ẽ': 'e', 'ẹ': 'e',
        'ê': 'e', 'ế': 'e', 'ề': 'e', 'ể': 'e', 'ễ': 'e', 'ệ': 'e',
        'ì': 'i', 'í': 'i', 'ỉ': 'i', 'ĩ': 'i', 'ị': 'i',
        'ò': 'o', 'ó': 'o', 'ỏ': 'o', 'õ': 'o', 'ọ': 'o',
        'ô': 'o', 'ố': 'o', 'ồ': 'o', 'ổ': 'o', 'ỗ': 'o', 'ộ': 'o',
        'ơ': 'o', 'ớ': 'o', 'ờ': 'o', 'ở': 'o', 'ỡ': 'o', 'ợ': 'o',
        'ù': 'u', 'ú': 'u', 'ủ': 'u', 'ũ': 'u', 'ụ': 'u',
        'ư': 'u', 'ứ': 'u', 'ừ': 'u', 'ử': 'u', 'ữ': 'u', 'ự': 'u',
        'ỳ': 'y', 'ý': 'y', 'ỷ': 'y', 'ỹ': 'y', 'ỵ': 'y',
    }
    
    result = ''
    for char in text:
        if char in vietnamese_map:
            result += vietnamese_map[char]
        elif char.isascii() and char.isalnum():
            result += char
        elif char in (' ', ',', '-', '(', ')', '/'):
            result += '_'
        # Skip other characters
    
    # Clean up underscores
    result = re.sub(r'_+', '_', result)
    result = result.strip('_')
    
    return result


def safe_float(value, default=0.0):
    """Safely convert value to float"""
    if value is None or value == '' or value == '-':
        return default
    try:
        return float(value)
    except (ValueError, TypeError):
        return default


def safe_int(value, default=0):
    """Safely convert value to int"""
    if value is None or value == '' or value == '-':
        return default
    try:
        return int(float(value))
    except (ValueError, TypeError):
        return default


def main():
    # Load Excel
    wb = openpyxl.load_workbook('assets/Bảng tính calo thực phẩm.xlsx')
    ws = wb['TPTP 2017']
    
    print(f"Reading sheet 'TPTP 2017': {ws.max_row} rows, {ws.max_column} cols")
    
    foods = []
    seen_keys = set()
    skipped = 0
    
    # Data starts from row 2 (row 1 is header)
    for row_idx in range(2, ws.max_row + 1):
        name = ws.cell(row=row_idx, column=2).value
        
        # Skip empty rows
        if not name or str(name).strip() == '':
            skipped += 1
            continue
        
        name = str(name).strip()
        calories = safe_float(ws.cell(row=row_idx, column=4).value)
        protein = safe_float(ws.cell(row=row_idx, column=5).value)
        fat = safe_float(ws.cell(row=row_idx, column=6).value)
        carbs = safe_float(ws.cell(row=row_idx, column=7).value)
        fiber = safe_float(ws.cell(row=row_idx, column=8).value)
        gi = safe_int(ws.cell(row=row_idx, column=10).value, default=50)
        
        # Tạo key unique
        key = create_key(name)
        if key in seen_keys:
            # Append số để tránh trùng
            counter = 2
            while f"{key}_{counter}" in seen_keys:
                counter += 1
            key = f"{key}_{counter}"
        seen_keys.add(key)
        
        food_item = {
            "key": key,
            "name": name,
            "name_en": "",
            "calories_per_100g": calories,
            "glycemic_index": gi,
            "protein": protein,
            "carbs": carbs,
            "fat": fat,
            "fiber": fiber,
            "category": "Chưa phân loại"
        }
        
        foods.append(food_item)
    
    # Write JSON output
    output_path = 'assets/database/food_database_from_excel.json'
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(foods, f, ensure_ascii=False, indent=2)
    
    print(f"\n✅ Đã convert thành công!")
    print(f"   Tổng số thực phẩm: {len(foods)}")
    print(f"   Bỏ qua (dòng trống): {skipped}")
    print(f"   Output: {output_path}")
    
    # In 5 mẫu đầu
    print(f"\n📋 5 mẫu đầu tiên:")
    for item in foods[:5]:
        print(f"   - {item['name']} ({item['calories_per_100g']} kcal, GI={item['glycemic_index']})")


if __name__ == '__main__':
    main()
